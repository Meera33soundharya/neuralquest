import docx

doc = docx.Document()

# Title
doc.add_heading("NeuralDocs — End-to-End Enterprise RAG Architecture & Manual", level=0)

# Section 1
doc.add_heading("1. Executive Summary", level=1)
doc.add_paragraph(
    "This document outlines the complete end-to-end design, implementation, and operational "
    "guidelines for the NeuralDocs Retrieval-Augmented Generation (RAG) platform. NeuralDocs enables "
    "organizations to transform static documentation into an interactive, high-accuracy AI knowledge assistant."
)

# Section 2
doc.add_heading("2. System Architecture & Core Layers", level=1)
doc.add_paragraph(
    "The NeuralDocs pipeline operates across six distinct operational layers:"
)

components = [
    ("Document Ingestion Layer", "Accepts PDF documents, validates file types, and extracts structured text alongside page numbers using PyMuPDF."),
    ("Chunking & Normalization", "Segments lengthy documents into semantic chunks (500 tokens) with sliding window overlap (50 tokens) to preserve conversational context across boundaries."),
    ("Embedding Generation", "Transforms extracted text segments into dense vector representations using state-of-the-art embedding models (FastEmbed / Google Gemini embeddings)."),
    ("Vector Storage & Indexing", "Stores vectors and metadata inside ChromaDB with HNSW indexing for low-latency similarity queries."),
    ("Context Retrieval & Reranking", "Vectorizes incoming user queries, executes cosine distance nearest-neighbor queries, and filters chunks isolated by user permissions."),
    ("LLM Generation & Synthesis", "Constructs grounded system prompts containing top-ranked context chunks and dispatches them to Google Gemini 1.5 to formulate sourced answers.")
]

for name, desc in components:
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(f"{name}: ")
    run.bold = True
    p.add_run(desc)

# Section 3
doc.add_heading("3. Security, Isolation & Multi-Tenancy", level=1)
doc.add_paragraph(
    "Every document and conversation is tied strictly to the authenticated user ID. When querying ChromaDB, "
    "metadata filters ensure that tenant A cannot access, retrieve, or infer data belonging to tenant B."
)

# Section 4
doc.add_heading("4. Step-by-Step User Instructions", level=1)
steps = [
    "Access the web portal at http://localhost:8000/",
    "Log in using your registered credentials (e.g. admin / admin123).",
    "Click the 'Upload PDF' button located on the left navigation panel.",
    "Wait for the animated progress bar to confirm embedding and indexing completion.",
    "Type natural language inquiries in the chat box and receive instant, referenced answers."
]
for i, step in enumerate(steps, 1):
    doc.add_paragraph(f"{i}. {step}")

# Section 5
doc.add_heading("5. Verification & Benchmark Questions", level=1)
doc.add_paragraph("Operators can ask the chatbot the following benchmark questions to verify system performance:")
questions = [
    "What chunking size and overlap window does NeuralDocs utilize?",
    "Which vector index is used inside ChromaDB for similarity queries?",
    "How does NeuralDocs ensure multi-tenant security and document isolation?",
    "What are the six operational layers of the RAG pipeline?"
]
for q in questions:
    doc.add_paragraph(q, style="List Bullet")

output_path = "NeuralDocs_End_to_End_Guide.docx"
doc.save(output_path)
print(f"Document saved to {output_path}")
